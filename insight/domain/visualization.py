from __future__ import annotations

from typing import Any, Literal, Optional

import pandas as pd
from pydantic import BaseModel, Field

from insight.domain.execution import DataTable


class PlotArtifact(BaseModel):
    path: str
    title: str
    chart_type: str
    step_id: Optional[int] = None


class ValidationIssue(BaseModel):
    severity: Literal["error", "warning"]
    message: str
    step_id: Optional[int] = None


class ValidationResult(BaseModel):
    valid: bool
    errors: list[ValidationIssue] = Field(default_factory=list)
    warnings: list[ValidationIssue] = Field(default_factory=list)
    checks_run: list[str] = Field(default_factory=list)

    def error_messages(self) -> list[str]:
        return [i.message for i in self.errors]

    def warning_messages(self) -> list[str]:
        return [i.message for i in self.warnings]


class Evidence(BaseModel):
    kind: Literal["calculation", "table", "warning"]
    label: str
    value: Optional[Any] = None
    detail: Optional[str] = None


class AnalysisResponse(BaseModel):
    question: str
    answer: str
    evidence: list[Evidence] = Field(default_factory=list)
    tables: list[DataTable] = Field(default_factory=list)
    plots: list[PlotArtifact] = Field(default_factory=list)
    validation: Optional[ValidationResult] = None
    meta: dict[str, Any] = Field(default_factory=dict)
    error: Optional[str] = None


def build_findings_section(doc, findings: list[dict]) -> None:
    """Render the zero-prompt briefing as the opening section of a report."""
    from docx.shared import Pt

    if not findings:
        return
    doc.add_heading("Top Findings", level=1)
    for finding in findings:
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(finding.get("title", ""))
        run.bold = True
        if finding.get("detail"):
            doc.add_paragraph(finding["detail"], style="List Bullet 2")
        if finding.get("suggested_query"):
            q = doc.add_paragraph(f"Suggested investigation: {finding['suggested_query']}")
            q.runs[0].font.size = Pt(9)
            q.runs[0].italic = True
