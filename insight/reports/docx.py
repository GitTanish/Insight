from __future__ import annotations

import io
from pathlib import Path


def build_report_docx(
    turns: list[dict],
    dataset_name: str,
    date_str: str,
    findings: list[dict] | None = None,
) -> bytes:
    """Render the full analysis session (briefing + turns) as a .docx report."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    from insight.domain.visualization import build_findings_section

    doc = Document()
    title = f"Insight Analysis Report — {dataset_name}"
    if date_str:
        title += f" — {date_str}"
    doc.add_heading(title, 0)

    build_findings_section(doc, findings or [])

    figure_count = 1
    for turn in turns:
        response = turn["response"]

        doc.add_heading(f"Q: {turn['question']}", level=1)
        doc.add_paragraph(response.answer)

        for plot in response.plots:
            path = Path(plot.path)
            if path.exists():
                try:
                    doc.add_picture(str(path), width=Inches(6.0))
                    caption = doc.add_paragraph(f"Figure {figure_count}: {plot.title}")
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if caption.runs:
                        caption.runs[0].italic = True
                        caption.runs[0].font.size = Pt(9)
                    figure_count += 1
                except Exception:
                    doc.add_paragraph(f"[figure unavailable: {plot.title}]")

        calc_evidence = [e for e in response.evidence if e.kind == "calculation"]
        warning_evidence = [e for e in response.evidence if e.kind == "warning"]

        if calc_evidence:
            doc.add_heading("Evidence", level=2)
            for item in calc_evidence:
                line = f"- {item.label}: {item.value}"
                if item.detail:
                    line += f" ({item.detail})"
                doc.add_paragraph(line, style="List Bullet")

        if warning_evidence:
            doc.add_heading("Caveats", level=2)
            for item in warning_evidence:
                doc.add_paragraph(item.label, style="List Bullet")

        if response.meta.get("model"):
            footer = doc.add_paragraph(f"[analysis by {response.meta['model']}]")
            if footer.runs:
                footer.runs[0].font.size = Pt(8)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
